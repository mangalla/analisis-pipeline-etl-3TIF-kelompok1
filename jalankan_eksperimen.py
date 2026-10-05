"""EKSPERIMEN PERFORMA PIPELINE ETL
Projek Statistika & Probabilitas - Kelas 3TIF, Kelompok 01

Cara pakai (dari folder ini):
    python jalankan_eksperimen.py --uji     -> uji coba singkat (6 run, +-1 menit), tidak masuk data
    python jalankan_eksperimen.py           -> eksperimen penuh (30 pemanasan + 300 run), lalu membuat Excel
    python jalankan_eksperimen.py --excel   -> hanya membuat ulang file Excel dari hasil_mentah.csv

Kalau eksperimen terhenti di tengah jalan (laptop mati, tertutup, dsb.), cukup jalankan
perintah yang sama lagi: run yang sudah tercatat tidak diulang, eksperimen dilanjutkan.
Untuk mengulang dari nol: python jalankan_eksperimen.py --baru

Desain: 3 format file x 5 volume data x 2 library x 10 ulangan = 300 run.
"""
import argparse
import csv
import datetime as dt
import itertools
import json
import os
import platform
import random
import subprocess
import sys
import time

KELAS, KELOMPOK = "3TIF", "01"
NAMA_FILE_EXCEL = f"{KELAS}_{KELOMPOK}_PipelineETL.xlsx"
JUDUL = ("Analisis Pengaruh Format File dan Volume Data terhadap Waktu Eksekusi "
         "dan Penggunaan Memori Pipeline ETL Berbasis Python")

FORMATS = ["CSV", "JSON", "Parquet"]
EXT = {"CSV": "csv", "JSON": "jsonl", "Parquet": "parquet"}
VOLUMES = [10_000, 50_000, 100_000, 500_000, 1_000_000]
LIBRARIES = ["Pandas", "Polars"]
ULANGAN = 10
SEED = 42
MAKS_COBA = 3  # run yang gagal dicoba ulang maksimal 3 kali sebelum ditandai NA

HERE = os.path.dirname(os.path.abspath(__file__))
DIR_INPUT = os.path.join(HERE, "data_input")
DIR_OUTPUT = os.path.join(HERE, "data_output")
F_IDENTITAS = os.path.join(HERE, "identitas_eksperimen.json")
F_PEMANASAN = os.path.join(HERE, "log_pemanasan.csv")
F_MENTAH = os.path.join(HERE, "hasil_mentah.csv")
F_LOG = os.path.join(HERE, "log_eksperimen.txt")

KOLOM_MENTAH = ["urutan_eksekusi", "waktu_mulai_run", "format_file", "library", "jumlah_baris",
                "ulangan_ke", "ukuran_file_mb", "waktu_baca_detik", "waktu_transformasi_detik",
                "waktu_tulis_detik", "total_waktu_detik", "memori_puncak_mb",
                "throughput_baris_per_detik", "status_run", "jumlah_percobaan"]


# ----------------------------------------------------------------------------- util
def log(pesan):
    baris = f"[{dt.datetime.now():%Y-%m-%d %H:%M:%S}] {pesan}"
    print(baris, flush=True)
    with open(F_LOG, "a", encoding="utf-8") as f:
        f.write(baris + "\n")


def cek_library():
    kurang = []
    for nama in ["numpy", "pandas", "polars", "pyarrow", "psutil", "openpyxl"]:
        try:
            __import__(nama)
        except ImportError:
            kurang.append(nama)
    if kurang:
        print("\nLibrary berikut belum terpasang: " + ", ".join(kurang))
        print("Pasang dulu dengan perintah:\n")
        print("    python -m pip install -r requirements.txt\n")
        sys.exit(1)


def nama_cpu():
    sistem = platform.system()
    try:
        if sistem == "Windows":
            import winreg
            k = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE,
                               r"HARDWARE\DESCRIPTION\System\CentralProcessor\0")
            return winreg.QueryValueEx(k, "ProcessorNameString")[0].strip()
        if sistem == "Linux":
            with open("/proc/cpuinfo") as f:
                for line in f:
                    if line.startswith("model name"):
                        return line.split(":", 1)[1].strip()
        if sistem == "Darwin":
            return subprocess.check_output(
                ["sysctl", "-n", "machdep.cpu.brand_string"], text=True).strip()
    except Exception:
        pass
    return platform.processor() or "tidak terdeteksi"


def status_daya():
    import psutil
    try:
        b = psutil.sensors_battery()
    except Exception:
        b = None
    if b is None:
        return "tidak ada baterai / tidak terdeteksi"
    return f"{'terhubung charger' if b.power_plugged else 'TIDAK terhubung charger'} (baterai {b.percent:.0f}%)"


def info_lingkungan():
    import numpy, pandas, polars, pyarrow, psutil
    return {
        "os": f"{platform.system()} {platform.release()} ({platform.version()})",
        "cpu": nama_cpu(),
        "cpu_inti_fisik": psutil.cpu_count(logical=False),
        "cpu_inti_logis": psutil.cpu_count(logical=True),
        "ram_total_gb": round(psutil.virtual_memory().total / 1024 ** 3, 1),
        "python": platform.python_version(),
        "numpy": numpy.__version__,
        "pandas": pandas.__version__,
        "polars": polars.__version__,
        "pyarrow": pyarrow.__version__,
    }


def tanya(pertanyaan, wajib=True):
    while True:
        jawab = input(pertanyaan).strip()
        if jawab or not wajib:
            return jawab
        print("  (wajib diisi)")


def siapkan_identitas(baru):
    if os.path.exists(F_IDENTITAS) and not baru:
        with open(F_IDENTITAS, encoding="utf-8") as f:
            return json.load(f)
    print("\n=== IDENTITAS EKSPERIMEN (dicatat otomatis ke file Excel) ===")
    ident = {
        "operator": tanya("Nama anggota yang menjalankan eksperimen : "),
        "anggota_hadir": tanya("Nama anggota kelompok yang hadir (pisahkan koma): "),
        "laptop": tanya("Merek & tipe laptop (mis. ASUS Vivobook 14 A1404) : "),
        "lokasi": tanya("Lokasi eksperimen (mis. Lab PCR / kos) : "),
    }
    with open(F_IDENTITAS, "w", encoding="utf-8") as f:
        json.dump(ident, f, indent=2, ensure_ascii=False)
    return ident


# ----------------------------------------------------------------------------- data input
def buat_data_sintetis(n):
    """Data transaksi sintetis (bukan data pribadi siapa pun) sebagai BEBAN UJI pipeline."""
    import numpy as np
    import pandas as pd
    rng = np.random.default_rng(SEED)
    df = pd.DataFrame({
        "id_transaksi": np.arange(1, n + 1),
        "tanggal": pd.to_datetime("2026-01-01") + pd.to_timedelta(rng.integers(0, 270, n), unit="D"),
        "id_pelanggan": rng.integers(1, max(n // 10, 100), n),
        "kategori": rng.choice(["Elektronik", "Pakaian", "Makanan", "Buku", "Olahraga"], n),
        "jumlah": rng.integers(1, 20, n),
        "harga_satuan": np.round(rng.uniform(5_000, 500_000, n), 2),
        "metode_bayar": rng.choice(["Tunai", "Transfer", "E-Wallet", "Kartu"], n),
        "status": rng.choice(["SELESAI", "BATAL", "PROSES"], n, p=[0.8, 0.1, 0.1]),
    })
    df["tanggal"] = df["tanggal"].dt.strftime("%Y-%m-%d")
    return df


def siapkan_input(volumes):
    os.makedirs(DIR_INPUT, exist_ok=True)
    os.makedirs(DIR_OUTPUT, exist_ok=True)
    for n in volumes:
        paths = {fmt: os.path.join(DIR_INPUT, f"data_{n}.{EXT[fmt]}") for fmt in FORMATS}
        if all(os.path.exists(p) for p in paths.values()):
            continue
        log(f"Membuat data input {n:,} baris (CSV, JSON, Parquet)...")
        df = buat_data_sintetis(n)
        df.to_csv(paths["CSV"], index=False)
        df.to_json(paths["JSON"], orient="records", lines=True)
        df.to_parquet(paths["Parquet"], index=False)


# ----------------------------------------------------------------------------- satu run
def satu_run(fmt, n, lib):
    src = os.path.join(DIR_INPUT, f"data_{n}.{EXT[fmt]}")
    dst = os.path.join(DIR_OUTPUT, f"hasil_{lib}_{n}.{EXT[fmt]}")
    ukuran = os.path.getsize(src) / 1024 / 1024
    for coba in range(1, MAKS_COBA + 1):
        p = subprocess.run([sys.executable, os.path.join(HERE, "run_one.py"), lib, fmt, src, dst],
                           capture_output=True, text=True)
        if p.returncode == 0:
            r = json.loads(p.stdout.strip().splitlines()[-1])
            total = r["waktu_baca"] + r["waktu_transformasi"] + r["waktu_tulis"]
            return {"ukuran_file_mb": ukuran, "waktu_baca_detik": r["waktu_baca"],
                    "waktu_transformasi_detik": r["waktu_transformasi"],
                    "waktu_tulis_detik": r["waktu_tulis"], "total_waktu_detik": total,
                    "memori_puncak_mb": r["memori_puncak_mb"],
                    "throughput_baris_per_detik": n / total,
                    "status_run": "OK", "jumlah_percobaan": coba}
        log(f"  Run GAGAL (percobaan {coba}) {fmt} {lib} {n}: {p.stderr.strip()[-300:]}")
    kosong = {k: "NA" for k in ["waktu_baca_detik", "waktu_transformasi_detik", "waktu_tulis_detik",
                                "total_waktu_detik", "memori_puncak_mb", "throughput_baris_per_detik"]}
    return {"ukuran_file_mb": ukuran, **kosong, "status_run": "GAGAL", "jumlah_percobaan": MAKS_COBA}


def fmt_durasi(detik):
    m, s = divmod(int(detik), 60)
    return f"{m} menit {s} detik" if m else f"{s} detik"


# ----------------------------------------------------------------------------- mode uji
def mode_uji():
    log("=== MODE UJI: 6 run singkat (10.000 baris), hasil TIDAK masuk data ===")
    siapkan_input([10_000])
    print(f"\n{'format':<8} {'library':<7} {'total_waktu(s)':>15} {'memori(MB)':>11}")
    for fmt, lib in itertools.product(FORMATS, LIBRARIES):
        r = satu_run(fmt, 10_000, lib)
        if r["status_run"] == "OK":
            print(f"{fmt:<8} {lib:<7} {r['total_waktu_detik']:>15.4f} {r['memori_puncak_mb']:>11.2f}")
        else:
            print(f"{fmt:<8} {lib:<7} GAGAL - lihat log_eksperimen.txt")
    print("\nKalau keenam baris berisi angka, laptop SIAP untuk eksperimen penuh:")
    print("    python jalankan_eksperimen.py\n")


# ----------------------------------------------------------------------------- eksperimen penuh
def daftar_run():
    kombinasi = [(fmt, n, lib, u) for fmt, n, lib in itertools.product(FORMATS, VOLUMES, LIBRARIES)
                 for u in range(1, ULANGAN + 1)]
    random.Random(SEED).shuffle(kombinasi)  # urutan acak tapi tetap sama setiap kali (bisa dilanjutkan)
    return kombinasi


def baca_csv(path):
    """Baca hasil yang sudah tercatat. Baris terakhir yang terpotong (laptop mati saat menulis)
    dibuang supaya run tersebut diulang dengan benar."""
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        kolom = reader.fieldnames or []
        rows = list(reader)
    utuh = [r for r in rows if None not in r and r.get(kolom[-1]) not in (None, "")] if kolom else []
    if len(utuh) != len(rows) or not kolom:
        with open(path, "w", newline="", encoding="utf-8") as f:
            if kolom:
                w = csv.DictWriter(f, fieldnames=kolom)
                w.writeheader()
                w.writerows(utuh)
        log(f"Membuang {len(rows) - len(utuh)} baris tidak utuh di {os.path.basename(path)}")
        if not kolom:
            os.remove(path)
    return utuh


def eksperimen_penuh(baru):
    if baru:
        for p in [F_PEMANASAN, F_MENTAH, F_IDENTITAS]:
            if os.path.exists(p):
                os.remove(p)
    ident = siapkan_identitas(baru)
    env = info_lingkungan()
    with open(os.path.join(HERE, "lingkungan.json"), "w", encoding="utf-8") as f:
        json.dump(env, f, indent=2)

    log("=== MULAI EKSPERIMEN PENUH ===")
    log(f"Operator: {ident['operator']} | Laptop: {ident['laptop']} | CPU: {env['cpu']} | "
        f"RAM: {env['ram_total_gb']} GB | {env['os']}")
    daya = status_daya()
    log(f"Status daya saat mulai: {daya}")
    if "TIDAK" in daya:
        print("\n!! Laptop TIDAK terhubung charger. Colokkan charger agar hasil tidak dipengaruhi mode hemat daya.")
        if input("   Lanjutkan saja? (y/n): ").strip().lower() != "y":
            sys.exit(0)

    siapkan_input(VOLUMES)

    # ---- tahap 1: pemanasan (1x per kombinasi, tidak masuk data)
    sudah_panas = baca_csv(F_PEMANASAN)
    if len(sudah_panas) < 30:
        log("Tahap 1/2: PEMANASAN 30 run (tidak masuk data, mengurangi efek run pertama/cache dingin)")
        baru_file = not os.path.exists(F_PEMANASAN)
        with open(F_PEMANASAN, "a", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            if baru_file:
                w.writerow(["no", "format_file", "library", "jumlah_baris", "total_waktu_detik", "status_run"])
            for i, (fmt, n, lib) in enumerate(itertools.product(FORMATS, VOLUMES, LIBRARIES), start=1):
                if i <= len(sudah_panas):
                    continue
                r = satu_run(fmt, n, lib)
                w.writerow([i, fmt, lib, n, r["total_waktu_detik"], r["status_run"]])
                f.flush()
                print(f"  pemanasan {i}/30", end="\r", flush=True)
        print()
    else:
        log("Tahap 1/2: pemanasan sudah selesai sebelumnya, dilewati.")

    # ---- tahap 2: 300 run utama
    runs = daftar_run()
    sudah = baca_csv(F_MENTAH)
    mulai_dari = len(sudah)
    if mulai_dari >= len(runs):
        log("Tahap 2/2: semua 300 run sudah tercatat.")
    else:
        if mulai_dari:
            log(f"Tahap 2/2: MELANJUTKAN dari run ke-{mulai_dari + 1} (run 1-{mulai_dari} sudah tercatat)")
            log(f"Status daya saat dilanjutkan: {status_daya()}")
        else:
            log("Tahap 2/2: 300 RUN UTAMA (urutan acak, seed 42)")
        baru_file = not os.path.exists(F_MENTAH)
        t_mulai = time.time()
        with open(F_MENTAH, "a", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=KOLOM_MENTAH)
            if baru_file:
                w.writeheader()
            for urutan, (fmt, n, lib, u) in enumerate(runs, start=1):
                if urutan <= mulai_dari:
                    continue
                stempel = f"{dt.datetime.now():%Y-%m-%d %H:%M:%S}"
                r = satu_run(fmt, n, lib)
                w.writerow({"urutan_eksekusi": urutan, "waktu_mulai_run": stempel, "format_file": fmt,
                            "library": lib, "jumlah_baris": n, "ulangan_ke": u, **r})
                f.flush()
                selesai = urutan - mulai_dari
                sisa = (time.time() - t_mulai) / selesai * (len(runs) - urutan)
                print(f"  run {urutan}/300  ({fmt:<7} {lib:<6} {n:>9,} baris)  "
                      f"perkiraan sisa: {fmt_durasi(sisa):<20}", end="\r", flush=True)
                if urutan % 50 == 0:
                    print()
                    log(f"  {urutan}/300 run tercatat")
        print()
    log(f"Status daya saat selesai: {status_daya()}")
    log("=== EKSPERIMEN SELESAI ===")
    buat_excel()


# ----------------------------------------------------------------------------- excel
def buat_excel():
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    mentah = pd.read_csv(F_MENTAH, na_values=["NA"], keep_default_na=False)
    if len(mentah) < 300:
        log(f"PERINGATAN: baru {len(mentah)} dari 300 run tercatat. Jalankan lagi untuk melanjutkan.")
    with open(F_IDENTITAS, encoding="utf-8") as f:
        ident = json.load(f)
    with open(os.path.join(HERE, "lingkungan.json"), encoding="utf-8") as f:
        env = json.load(f)

    d = mentah.copy()
    d["format_file"] = pd.Categorical(d["format_file"], FORMATS, ordered=True)
    d = d.sort_values(["format_file", "jumlah_baris", "library", "ulangan_ke"]).reset_index(drop=True)
    d["format_file"] = d["format_file"].astype(str)
    d.insert(0, "id_run", [f"R{i:03d}" for i in range(1, len(d) + 1)])
    kolom = ["id_run", "urutan_eksekusi", "ulangan_ke", "format_file", "library", "jumlah_baris",
             "ukuran_file_mb", "waktu_baca_detik", "waktu_transformasi_detik", "waktu_tulis_detik",
             "total_waktu_detik", "memori_puncak_mb", "throughput_baris_per_detik"]
    d = d[kolom]
    jumlah_na = int(d["total_waktu_detik"].isna().sum())
    gagal_awal = int((mentah["jumlah_percobaan"] > 1).sum())

    waktu = pd.to_datetime(mentah["waktu_mulai_run"])
    t_awal, t_akhir = waktu.min(), waktu.max()

    FONT = "Arial"
    hdr_font = Font(name=FONT, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    body, bold = Font(name=FONT), Font(name=FONT, bold=True)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    angka = {"urutan_eksekusi": "0", "ulangan_ke": "0", "jumlah_baris": "#,##0",
             "ukuran_file_mb": "0.0000", "waktu_baca_detik": "0.000000",
             "waktu_transformasi_detik": "0.000000", "waktu_tulis_detik": "0.000000",
             "total_waktu_detik": "0.000000", "memori_puncak_mb": "0.000",
             "throughput_baris_per_detik": "#,##0.00"}
    bulat = {"ukuran_file_mb": 4, "waktu_baca_detik": 6, "waktu_transformasi_detik": 6,
             "waktu_tulis_detik": 6, "total_waktu_detik": 6, "memori_puncak_mb": 3,
             "throughput_baris_per_detik": 2}

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(kolom)
    for row in d.itertuples(index=False):
        nilai = []
        for c, v in zip(kolom, row):
            if isinstance(v, float) and pd.isna(v):
                nilai.append("NA")
            elif c in bulat:
                nilai.append(round(float(v), bulat[c]))
            elif c in ("urutan_eksekusi", "ulangan_ke", "jumlah_baris"):
                nilai.append(int(v))
            else:
                nilai.append(v)
        ws.append(nilai)
    for j, c in enumerate(kolom, start=1):
        h = ws.cell(row=1, column=j)
        h.font, h.fill, h.border = hdr_font, hdr_fill, box
        h.alignment = Alignment(horizontal="center")
        for i in range(2, len(d) + 2):
            x = ws.cell(row=i, column=j)
            x.font = body
            if c in angka:
                x.number_format = angka[c]
        ws.column_dimensions[get_column_letter(j)].width = max(len(c) + 3, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{len(d) + 1}"

    kv = wb.create_sheet("Keterangan_Variabel")
    kv_rows = [
        ("No", "Nama Variabel", "Keterangan", "Tipe Data", "Jenis", "Skala", "Satuan", "Peran"),
        (1, "id_run", "Kode unik setiap run eksperimen (R001-R300), diurutkan per kombinasi perlakuan", "Kualitatif", "Identitas", "Nominal", "-", "Identitas"),
        (2, "urutan_eksekusi", "Urutan run saat benar-benar dijalankan (hasil pengacakan, 1-300)", "Kuantitatif", "Diskrit", "Ordinal", "-", "Kontrol"),
        (3, "ulangan_ke", "Nomor ulangan dalam satu kombinasi perlakuan (1-10)", "Kuantitatif", "Diskrit", "Ordinal", "-", "Kontrol"),
        (4, "format_file", "Format file input dan output pipeline: CSV, JSON (newline-delimited), Parquet", "Kualitatif", "Kategorik", "Nominal", "-", "Faktor"),
        (5, "library", "Library pemrosesan data: Pandas atau Polars", "Kualitatif", "Kategorik", "Nominal", "-", "Faktor"),
        (6, "jumlah_baris", "Volume data input: 10.000; 50.000; 100.000; 500.000; 1.000.000 baris", "Kuantitatif", "Diskrit", "Rasio", "baris", "Faktor"),
        (7, "ukuran_file_mb", "Ukuran file input di disk", "Kuantitatif", "Kontinu", "Rasio", "MB", "Pendukung"),
        (8, "waktu_baca_detik", "Waktu tahap Extract (membaca file input)", "Kuantitatif", "Kontinu", "Rasio", "detik", "Pendukung"),
        (9, "waktu_transformasi_detik", "Waktu tahap Transform (filter status, hitung total_harga, agregasi per kategori, join)", "Kuantitatif", "Kontinu", "Rasio", "detik", "Pendukung"),
        (10, "waktu_tulis_detik", "Waktu tahap Load (menulis hasil ke format yang sama)", "Kuantitatif", "Kontinu", "Rasio", "detik", "Pendukung"),
        (11, "total_waktu_detik", "Total waktu eksekusi = waktu_baca + waktu_transformasi + waktu_tulis", "Kuantitatif", "Kontinu", "Rasio", "detik", "VARIABEL UTAMA 1"),
        (12, "memori_puncak_mb", "Kenaikan puncak memori proses selama ETL dibanding sebelum ETL dimulai", "Kuantitatif", "Kontinu", "Rasio", "MB", "VARIABEL UTAMA 2"),
        (13, "throughput_baris_per_detik", "Kecepatan pemrosesan = jumlah_baris / total_waktu_detik", "Kuantitatif", "Kontinu", "Rasio", "baris/detik", "VARIABEL UTAMA 3"),
    ]
    for r in kv_rows:
        kv.append(r)
    for j, w in enumerate([5, 28, 70, 13, 12, 10, 12, 18], start=1):
        kv.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(kv.iter_rows(min_row=1, max_row=len(kv_rows)), start=1):
        for cell in r:
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if i == 1:
                cell.font, cell.fill = hdr_font, hdr_fill
            else:
                utama = str(r[7].value).startswith("VARIABEL UTAMA")
                cell.font = bold if utama and cell.column in (2, 8) else body
    kv.freeze_panes = "A2"

    info = wb.create_sheet("Info_Dataset")
    info_rows = [
        ("Judul Projek", JUDUL),
        ("Kelas / Kelompok", f"{KELAS} / {KELOMPOK}"),
        ("Bidang", "Data Engineer"),
        ("Sumber Data", "mandiri (eksperimen)"),
        ("Teknik Pengumpulan Data",
         "Eksperimen terkontrol. Script Python (jalankan_eksperimen.py + run_one.py) menjalankan pipeline ETL "
         "dan mencatat waktu dengan time.perf_counter() serta puncak memori proses dengan psutil/OS. "
         "Setiap run dijalankan sebagai proses Python terpisah agar pengukuran memori tidak saling terpengaruh."),
        ("Unit Observasi", "Satu kali eksekusi (run) pipeline ETL"),
        ("Jumlah Observasi (Baris)", len(d)),
        ("Jumlah Variabel (Kolom)", len(kolom)),
        ("Desain Eksperimen",
         "Faktorial penuh 3 format file x 5 volume data x 2 library x 10 ulangan = 300 run. Didahului 30 run "
         "pemanasan yang tidak dicatat sebagai data. Urutan 300 run diacak (seed 42); urutan asli tersimpan "
         "di kolom urutan_eksekusi."),
        ("Data Input Pipeline",
         "Data transaksi sintetis dibangkitkan dengan seed tetap (42) sebagai beban uji: id_transaksi, tanggal, "
         "id_pelanggan, kategori, jumlah, harga_satuan, metode_bayar, status. Data yang dianalisis adalah hasil "
         "pengukuran performa, bukan data transaksi ini."),
        ("Tahapan Pipeline",
         "Extract: baca file input. Transform: filter status = SELESAI, hitung total_harga = jumlah x harga_satuan, "
         "agregasi total per kategori, join kembali ke data. Load: tulis hasil ke format yang sama."),
        ("Pelaksana Eksperimen", ident["operator"]),
        ("Anggota Hadir", ident["anggota_hadir"]),
        ("Lokasi", ident["lokasi"]),
        ("Laptop", ident["laptop"]),
        ("CPU", f"{env['cpu']} ({env['cpu_inti_fisik']} inti fisik, {env['cpu_inti_logis']} inti logis)"),
        ("RAM", f"{env['ram_total_gb']} GB"),
        ("Sistem Operasi", env["os"]),
        ("Versi Software", f"Python {env['python']}; numpy {env['numpy']}; pandas {env['pandas']}; "
                           f"polars {env['polars']}; pyarrow {env['pyarrow']}"),
        ("Waktu Pelaksanaan", f"{t_awal:%d-%m-%Y %H:%M} s.d. {t_akhir:%d-%m-%Y %H:%M} "
                              f"(stempel waktu tiap run ada di hasil_mentah.csv)"),
        ("Missing Value",
         f"{jumlah_na} run bernilai NA (gagal setelah {MAKS_COBA} kali percobaan). Tanda missing value: teks NA. "
         f"Run yang berhasil setelah dicoba ulang: {gagal_awal}."),
        ("Perubahan dari Data Asli",
         "Tidak ada data yang dihapus. Hanya pembulatan: waktu 6 desimal, memori 3 desimal, ukuran file 4 desimal, "
         "throughput 2 desimal. Data asli lengkap: hasil_mentah.csv dan log_eksperimen.txt."),
        ("Satuan", "Waktu dalam detik; memori dan ukuran file dalam MB (1 MB = 1.048.576 byte); "
                   "throughput dalam baris/detik."),
        ("Catatan Tambahan (isi manual)", "-"),
    ]
    for r in info_rows:
        info.append(r)
    info.column_dimensions["A"].width = 30
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows(min_row=1, max_row=len(info_rows)):
        row[0].font = bold
        row[0].fill = PatternFill("solid", fgColor="DDEBF7")
        row[1].font = body
        for cell in row:
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    out = os.path.join(HERE, NAMA_FILE_EXCEL)
    try:
        wb.save(out)
    except PermissionError:
        print(f"\nFile {NAMA_FILE_EXCEL} sedang terbuka di Excel. Tutup dulu, lalu jalankan:")
        print("    python jalankan_eksperimen.py --excel\n")
        sys.exit(1)
    log(f"File Excel dibuat: {NAMA_FILE_EXCEL} ({len(d)} baris x {len(kolom)} kolom, NA = {jumlah_na})")


# ----------------------------------------------------------------------------- main
if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Eksperimen performa pipeline ETL - 3TIF Kelompok 01")
    ap.add_argument("--uji", action="store_true", help="uji coba singkat 6 run")
    ap.add_argument("--baru", action="store_true", help="hapus hasil lama dan mulai dari nol")
    ap.add_argument("--excel", action="store_true", help="hanya buat ulang file Excel")
    a = ap.parse_args()
    os.chdir(HERE)
    cek_library()
    if a.uji:
        mode_uji()
    elif a.excel:
        buat_excel()
    else:
        try:
            eksperimen_penuh(a.baru)
        except KeyboardInterrupt:
            print()
            log("Eksperimen DIHENTIKAN (Ctrl+C). Jalankan perintah yang sama untuk melanjutkan.")
